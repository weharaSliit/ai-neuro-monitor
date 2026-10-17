"""Workbook of the simulation study (SIMULATED -- software-in-the-loop).

Values only (computed in Python): every number and PASS/FAIL is visible in any
viewer without recalculation. Sheets: README, Fidelity, Criteria, Results
(later occasion), Results (fresh ring calibration), Runs, Parameters,
Example sessions, Example calibrations.

  PYTHONPATH=../..:. ../../venv/bin/python -m simkit.workbook
"""
import json
from pathlib import Path

import numpy as np

from . import GENERATOR_VERSION
from .fidelity import fidelity
from .profiles import PROFILES, sample_user
from .summarize import summarize

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "output" / "simkit"
TAG = "SIMULATED – software-in-the-loop"
COMP = ("mind", "body", "fused")


def build(path=OUT / "mindpulse_simulation_study.xlsx", n_fid=24):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    F = "Arial"
    norm, bold, title = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=14, bold=True)
    note = Font(name=F, size=9, italic=True, color="666666")
    head = PatternFill("solid", fgColor="E8EEF7")
    ok_fill, bad_fill, part_fill = (PatternFill("solid", fgColor=c) for c in ("D9F2DD", "F8D7D3", "FFF0C2"))
    res = json.loads((OUT / "results.json").read_text())
    res_same = json.loads((OUT / "results_same_occasion.json").read_text())
    summ, summ_same = summarize(res), summarize(res_same, name="results_same_occasion.json")
    crit = json.loads((HERE / "criteria.json").read_text())
    meta = res["meta"]
    wb = Workbook()

    def sheet(name, heading, sub=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        ws["A1"], ws["A1"].font = f"{heading} · {TAG}", title
        if sub:
            ws["A2"], ws["A2"].font = sub, note
        return ws

    def table(ws, r0, cols, rows, widths=None, fmt=None):
        for c, h in enumerate(cols, 1):
            cell = ws.cell(row=r0, column=c, value=h)
            cell.font, cell.fill = bold, head
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for i, row in enumerate(rows, r0 + 1):
            for c, v in enumerate(row, 1):
                if isinstance(v, (np.floating, np.integer)):
                    v = v.item()
                if isinstance(v, (list, dict, tuple)):
                    v = json.dumps(v, default=float)
                cell = ws.cell(row=i, column=c, value=v)
                cell.font = norm
                if fmt and c - 1 in fmt and isinstance(v, (int, float)):
                    cell.number_format = fmt[c - 1]
                if v in ("PASS", "FAIL", "PARTIAL"):
                    cell.fill = {"PASS": ok_fill, "FAIL": bad_fill, "PARTIAL": part_fill}[v]
                    cell.font = bold
        for c, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
        return r0 + len(rows) + 2

    # ── README ──
    ws = sheet("README", "MindPulse simulation study", first=True)
    lines = [
        ("What", "Eight virtual-user profiles x 20 seeds, each running the full protocol. Their device data were processed by the SAME "
                 "backend code the deployed services run (imported read-only): EEG calibration, direction-aware ABA threshold, live "
                 "block scoring, ring calibration and prediction, fusion engine, 2-of-3 alerts. Nothing deployed was changed."),
        ("Signals", "EEG = real SAM40 T7-T8 trials (each virtual user is one SAM40 subject) + a headband layer measured on real "
                    "wearable02 good-contact blocks. Ring = IPFM beat model + skin conductance with Poisson SCRs, sized from real "
                    "ring windows (27 users, 1,038 windows) and WESAD class statistics."),
        ("Latent model", "Demand drives brain (onset 10-30 s / recovery 20-60 s), heart (15-40 / 60-120 s) and skin (20-60 / "
                         "120-300 s) with separate time constants per user, plus fatigue and unrelated arousal events."),
        ("Protocol", f"Settle 2 min -> rest 10 EEG blocks -> task 19 -> recovery 12 (5 min) -> live: calm 4 min, demand 4, calm 4, "
                     f"demand 4, calm 5. Ring windows (60 s): rest 4, task 7, recovery 5. Grace period: 2 min after demand ends is "
                     f"excluded from false alarms; first 2 min of demand excluded from the alert rate (onset)."),
        ("Main run", "Live monitoring on a LATER OCCASION than the ring calibration (ring baselines persist for days in the real "
                     "system): heart rate +-6 bpm, skin level x lognormal(0.35), variability x lognormal(0.25), sized from real "
                     "ring windows vs their stored baselines. A second run with a fresh ring calibration is the comparison."),
        ("Fidelity", "Simulated signals and model performance were checked against real readings BEFORE evaluation (Fidelity sheet). "
                     "23 of 24 checks pass; the ring model's output stays more extreme on simulated windows than on real ones (partial)."),
        ("Criteria", "Pass criteria and metric definitions were fixed before the evaluation (Criteria sheet, criteria.json)."),
        ("Limits", "Simulated users show how the system behaves on realistic signals; they are not evidence about real people. "
                   "The deployed EEG model was trained on all 40 SAM40 subjects, so EEG here is somewhat optimistic (in-sample). "
                   "Results belong in their own thesis section, separate from real-session results."),
        ("Generator", f"{GENERATOR_VERSION}; seeds {meta['seeds_per_profile']} per profile; run time {meta['runtime_s'] / 60:.1f} min."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=k).font = bold
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.alignment = norm, Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 48
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 16, 120

    # ── Fidelity ──
    ws = sheet("Fidelity", "Fidelity: simulated vs real", f"{n_fid} virtual users per check (typical profile). PASS = simulated median "
               "inside the real 5-95% range (+25% of its width); AUC targets as noted.")
    rows = []
    for r in fidelity(n_fid):
        s, ref = r["simulated"] or {}, r["real"] or {}
        verdict = "PARTIAL" if r["pass"] is False and "PARTIAL" in r.get("note", "") else ("PASS" if r["pass"] else "FAIL")
        rows.append([r["signal"], s.get("p5"), s.get("median"), s.get("p95"), r.get("mean"), ref.get("p5"), ref.get("median"),
                     ref.get("p95"), verdict, r.get("note", "")])
    table(ws, 4, ["Signal", "Sim p5", "Sim median", "Sim p95", "Sim mean", "Real p5", "Real median", "Real p95", "Result", "Note"],
          rows, [46, 10, 10, 10, 10, 10, 10, 10, 10, 90], {i: "0.000" for i in range(1, 8)})

    # ── Criteria ──
    ws = sheet("Criteria", "Pre-registered criteria and metric definitions", crit["_note"])
    r = table(ws, 4, ["Metric", "Definition"], [[k, v] for k, v in crit["metric_definitions"].items()], [34, 140])
    rows = [[p, m, op, v] for p, cs in crit["pass_criteria"].items() for m, (op, v) in cs.items()]
    table(ws, r, ["Profile", "Metric", "Test", "Target"], rows, [34, 34, 8, 10])

    # ── Results ──
    def results_sheet(name, heading, sub, S, rs):
        ws = sheet(name, heading, sub)
        rows = []
        for p, x in S.items():
            pr = PROFILES[p]
            rr = [q for q in rs["runs"] if q["profile"] == p and "error" not in q]
            row = [pr.title, pr.tests, x["n_runs"], json.dumps(x["eeg_status_counts"]), json.dumps(x["ring_status_counts"])]
            for k in COMP:
                for m in ("alert_rate_demand", "false_alarm_rate_calm"):
                    d = x[f"{k}_{m}"]
                    row += [d["mean"], d["lo"], d["hi"]]
            for k in COMP:
                d = x[f"{k}_time_to_alert_s"]
                row += [d["mean"], d.get("none_share")]
                row += [x[f"{k}_time_to_clear_s"]["mean"]]
            row += [x["gate_rejection_rate"]["mean"], x["w_eeg_mean"]["mean"], x["w_ring_mean"]["mean"],
                    x["fused_equal_alert_rate_demand"]["mean"], x["fused_equal_false_alarm_rate_calm"]["mean"],
                    float(np.mean([q.get("ring_rest_retries") or 0 for q in rr])) if rr else None]
            rows.append(row)
        cols = ["Profile", "Tests", "Runs", "EEG calibration (count)", "Ring calibration (count)"]
        for k in COMP:
            for m in ("alert in demand", "false alarm in calm"):
                cols += [f"{k.title()} {m}", "95% CI low", "95% CI high"]
        for k in COMP:
            cols += [f"{k.title()} time to alert (s)", f"{k.title()} no alert in a demand period (share)", f"{k.title()} time to clear (s)"]
        cols += ["EEG gate rejection", "Fusion weight EEG", "Fusion weight ring", "Equal-weight fused: alert in demand",
                 "Equal-weight fused: false alarm in calm", "Ring rest retries (mean)"]
        pct = {i: "0%" for i in range(5, 23)}
        pct.update({i: "0" for i in (23, 25, 26, 28, 29, 31)})
        pct.update({24: "0%", 27: "0%", 30: "0%", 32: "0%", 33: "0.00", 34: "0.00", 35: "0%", 36: "0%", 37: "0.0"})
        r = table(ws, 4, cols, rows, [26, 34, 6, 34, 34] + [11] * 32, pct)
        ws.cell(row=r, column=1, value="Pass criteria (fixed in advance)").font = bold
        crows = [[PROFILES[p].title, c["metric"], c["op"], c["target"], c["value"], "PASS" if c["pass"] else "FAIL"]
                 for p, x in S.items() for c in x["checks"]]
        table(ws, r + 1, ["Profile", "Metric", "Test", "Target", "Value", "Result"], crows, None, {3: "0.00", 4: "0.00"})
        ws.freeze_panes = "C5"

    results_sheet("Results (later occasion)", "Results: live monitoring on a later occasion (main)",
                  "Mean over 20 seeds per profile with 95% t-intervals. Rates exclude the 2 min onset of each demand period and the "
                  "2 min grace after it.", summ, res)
    results_sheet("Results (fresh ring cal)", "Comparison: live monitoring right after a fresh ring calibration",
                  "Same users and seeds; only the day-to-day ring shift is removed.", summ_same, res_same)

    # ── Runs ──
    ws = sheet("Runs", "Every run (later occasion)")
    keys = ["profile", "seed", "eeg_subject", "eeg_status", "eeg_direction", "eeg_separation_sd", "eeg_reversal_sd",
            "eeg_rest_blocks_seen", "eeg_amplitude_gain", "ring_status", "ring_rest_retries", "ring_limit",
            "mind_alert_rate_demand", "body_alert_rate_demand", "fused_alert_rate_demand",
            "mind_false_alarm_rate_calm", "body_false_alarm_rate_calm", "fused_false_alarm_rate_calm",
            "gate_rejection_rate", "w_eeg_mean", "w_ring_mean"]
    rows = [[q.get(k) for k in keys] for q in res["runs"] if "error" not in q]
    table(ws, 3, keys, rows, [20, 6, 9, 15, 8, 10, 10, 9, 9, 15, 8, 8] + [10] * 9,
          {5: "0.00", 6: "0.00", 8: "0.00", 11: "0.000", **{i: "0%" for i in range(12, 19)}, 19: "0.00", 20: "0.00"})

    # ── Parameters ──
    ws = sheet("Parameters", "Generator parameters", "Per-user values drawn from these ranges (profiles.py); the sampled values of every run follow.")
    rng_rows = [
        ["EEG base signal", "SAM40 T7-T8 trials of one subject per user (Relax = rest, Arithmetic/Stroop/Mirror = task)", "SAM40 (40 subjects)"],
        ["EEG subject strata", "typical: shift >= 0.5 spreads (30), strong >= 1.2 (11), weak |shift| < 0.4 (5), inverted <= -0.6 (4)", "reference/sam40_reference.json"],
        ["EEG headband RMS", "lognormal, median 11.4 uV, sd 0.25 (clip 6-30)", "real good-contact wearable02 blocks"],
        ["EEG drift / EMG", "drift 0.30-0.50 x RMS (0.3-3 Hz); muscle 0.10-0.22 x RMS (35-100 Hz), x3.5 bursts while tapping", "tuned to real delta/gamma share + AUC"],
        ["EEG blinks", "10-20 per min, 10-30 uV, 300 ms", "eyes-open rest"],
        ["Noise profile", "jaw clench 1/min (1-3 s), packet dropout 0.5/min, ring motion 1.5/min, GSR contact loss 0.4/min", "real-world artifacts"],
        ["Drowsy profile", "episodes every ~4 min in calm periods, 30-90 s: alpha -48%, theta up", ""],
        ["Drift profile", "after the task starts, alpha loss + slow theta growing over 10 min, never returns", "electrode / gel drift"],
        ["Fatigue", "0-0.3 per hour (all users), alpha down / theta up", ""],
        ["tau brain / heart / skin", "onset 10-30 / 15-40 / 20-60 s; recovery 20-60 / 60-120 / 120-300 s", "plan"],
        ["Arousal events", "0.12 per min, decay 15 s: HR +6 bpm, extra SCRs", "unrelated arousal"],
        ["HR baseline", "76-102 bpm (low-arousal profile -4)", "real ring users' median HR"],
        ["Ring stress effect", "HR +0.9-3.2 bpm, RSA -3-10 %, SCL +0.1-0.6 uS, SCR rate rest 1-3 -> +0.9-2.1 per min (x profile factor)",
         "sized for ring-model AUC ~0.8 (mild cognitive stressor)"],
        ["Beat jitter / errors", "30-55 ms timing jitter; 1.2 % missed, 0.6 % split beats", "real ring SD1 median 59 ms"],
        ["Slow variability", "HR wander +-4 bpm (tau 90 s); skin level +-6 % (tau 5 min); GSR noise 5 counts, level re-drawn each 30-90 s",
         "real ring windows vs stored baselines"],
        ["Later occasion", "HR +-6 bpm, SCL x lognormal(0.35), variability x lognormal(0.25)", "real ring windows vs stored baselines"],
    ]
    r = table(ws, 4, ["Parameter", "Range / value", "Source"], rng_rows, [24, 110, 44])
    prow = []
    for p in PROFILES:
        for s in range(meta["seeds_per_profile"]):
            up, _ = sample_user(p, 100 + s)
            e, g = up.eeg, up.ring
            prow.append([p, 100 + s, e["subject"], e["rms_uv"], e["drift_frac"], e["emg_frac"], g["hr0"], g["d_hr"], g["rsa_drop"],
                         g["scl0"], g["d_scl"], g["scr_rest"], g["scr_stress"], up.fatigue_per_hour,
                         *[x for m in ("brain", "heart", "skin") for x in up.taus[m]]])
    table(ws, r, ["Profile", "Seed", "SAM40 subject", "EEG RMS uV", "Drift frac", "EMG frac", "HR0", "dHR", "RSA drop", "SCL0",
                  "dSCL", "SCR rest/min", "SCR stress/min", "Fatigue/h", "tau brain on", "tau brain off", "tau heart on",
                  "tau heart off", "tau skin on", "tau skin off"], prow, None, {i: "0.00" for i in range(3, 20)})

    # ── Example sessions (seed 100 of each profile) ──
    ws = sheet("Example sessions", "Example live sessions (seed 100 of each profile)",
               "One row per 25 s block: latent loads, the three readings on the fusion scale (own limit 0.50) and alerts.")
    rows = []
    for q in res["runs"]:
        for lr in q.get("live_rows") or []:
            rows.append([q["profile"], lr["t"], lr["period"], lr["demand"], lr.get("load_brain"), lr.get("load_heart"), lr.get("load_skin"),
                         lr.get("mind"), lr.get("body"), lr.get("fused"), lr.get("fused_equal"), lr.get("mind_alert"), lr.get("body_alert"),
                         lr.get("fused_alert"), lr.get("w_eeg"), lr.get("w_ring"), lr.get("skipped"), lr.get("lost"), lr.get("drowsy"), lr.get("drift")])
    table(ws, 4, ["Profile", "Time (s)", "Period", "Demand", "Brain load", "Heart load", "Skin load", "Mind", "Body", "Fused",
                  "Fused (equal weights)", "Mind alert", "Body alert", "Fused alert", "EEG weight", "Ring weight", "EEG skipped",
                  "EEG block lost", "Drowsy", "Drift"], rows, [20, 9, 9, 8] + [10] * 16, {i: "0.00" for i in range(4, 20)})

    # ── Example calibrations ──
    ws = sheet("Example calibrations", "Example calibrations (seed 100 of each profile)", "Per-block EEG model scores and ring window scores.")
    rows = []
    for q in res["runs"]:
        c = q.get("calibration")
        if not c:
            continue
        for ph in ("eeg_rest", "eeg_task", "eeg_recovery", "ring_rest", "ring_task", "ring_recovery"):
            for i, v in enumerate(c.get(ph) or [], 1):
                rows.append([q["profile"], ph.split("_")[0].upper(), ph.split("_")[1], i, v,
                             c.get("eeg_baseline_logit") if ph.startswith("eeg") else None,
                             c.get("eeg_limit_logit") if ph.startswith("eeg") else q.get("ring_limit"), q.get(f"{ph.split('_')[0]}_status")])
    table(ws, 4, ["Profile", "Device", "Phase", "#", "Score (probability)", "EEG baseline (log-odds)", "Limit (EEG log-odds / ring prob)",
                  "Calibration result"], rows, [20, 8, 10, 5, 16, 16, 18, 18], {4: "0.00E+00", 5: "0.00", 6: "0.000"})

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())
