"""
One simulated session (the live stimulator's virtual person, on a simulated
clock) run through the REAL backend code, exported to Excel so every step can
be checked by hand:

  Summary          settings, backend results, Excel's own recomputation of
                   each decision (formulas), and PASS/FAIL checks
  EEG calibration  every 25 s block: phase, model score, log-odds, orientation,
                   fusion-scale value, above-limit flag (formulas)
  Ring calibration every 60 s window: phase, stress probability, above-limit flag
  Live monitoring  every 25 s block after calibration: demand, body states,
                   mind / body / fused (fusion engine), 2-of-3 alerts (formulas)

Same pipeline as live: live_stimulator.EEGGenerator/RingGenerator ->
parse_bipolar_payload (bridge) -> RealtimeInference (calibrate, task,
recovery, finalize_personal_threshold, infer_bipolar_block) ->
make_eeg_output_record / make_eeg_fusion_record -> fusion_engin; ring:
extract_features -> compute_baseline / predict_window ->
midpoint_stress_threshold (as calibration_capture.py).

  venv/bin/python backend/simulation/export_session_excel.py [--profile typical|inverted|drift]
"""
import argparse
import math
import sys
from pathlib import Path

import numpy as np

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "backend" / "simulation"))
sys.path.insert(0, str(REPO / "fushion engin"))

import live_stimulator as S  # noqa: E402
import fusion_engin  # noqa: E402
from backend.physiological.extract_features import compute_signal_quality, extract_features  # noqa: E402
from backend.physiological.predictor import (  # noqa: E402
    PERSONAL_THRESHOLD_MIN_SEPARATION, STRESS_DECISION_THRESHOLD, TASK_RAMP_SKIP_WINDOWS, compute_baseline,
    midpoint_stress_threshold, predict_window,
)
from backend.psychological.realtime.firebase_sync import make_eeg_fusion_record, make_eeg_output_record  # noqa: E402
from backend.psychological.realtime.mqtt_client import parse_bipolar_payload  # noqa: E402
from backend.psychological.realtime.pipeline import BLOCK_SEC, RealtimeInference  # noqa: E402
from backend.psychological.sam40.decision import (  # noqa: E402
    MIN_RECOVERY_REVERSAL_SD, MIN_TASK_RESPONSE_SD, TASK_RAMP_SKIP_BLOCKS,
)

STEP = S.EEG_BATCH / S.FS          # 0.2 s: one EEG batch and one ring packet per step
SETTLE, N_REST, N_TASK, N_REC = 120, 10, 19, 7
RING_REST, RING_TASK, RING_REC = 4, 8, 3
LIVE = S.LIVE_SCENARIO              # calm 240 s -> demand 300 s -> calm 300 s


class Session:
    """Streams the virtual person on a simulated clock, collecting EEG
    samples (as the bridge parses them) and ring packets."""

    def __init__(self, profile):
        self.p = S.Person(profile)
        self.eeg = S.EEGGenerator(self.p)
        self.ring = S.RingGenerator(self.p)
        self.t = 0.0

    def run(self, seconds, demand, tapping=None):
        p = self.p
        p.demand = demand
        p.tapping = demand > 0.5 if tapping is None else tapping
        samples, packets, loads = [], [], []
        for _ in range(int(round(seconds / STEP))):
            p.step(STEP)
            samples += parse_bipolar_payload(self.eeg.batch(S.EEG_BATCH))
            packets.append(self.ring.packet())
            self.t += STEP
        loads = (p.eeg_load, p.hr_load, p.eda_load)
        return samples, packets, loads


def ring_window(packets):
    """60 s of packets -> the arrays the ring bridge builds (|acc| in g)."""
    ibi = [b for d in packets for b in d["ibi"]]
    gsr = [d["gsr"] for d in packets]
    acc = [math.sqrt(d["acc_x"] ** 2 + d["acc_y"] ** 2 + d["acc_z"] ** 2) / 9.81 for d in packets]
    f2, motion, _fd, fallback, _wd = extract_features(ibi, gsr, acc, S.GSR_CAL, actual_fs=5)
    q = compute_signal_quality(ibi, gsr, acc, motion, actual_fs=5)
    return (None if motion or fallback else f2), q


def simulate(profile):
    s = Session(profile)
    eng = RealtimeInference(REPO / "backend" / "psychological" / "models")
    out = {"profile": profile, "eeg_cal": [], "ring_cal": [], "live": []}

    s.run(SETTLE, 0.0)                                   # settle: discarded by both devices
    # Rest: EEG 10 x 25 s (250 s); ring 4 x 60 s inside it.
    ring_packets = []
    for i in range(N_REST):
        smp, pk, _ = s.run(BLOCK_SEC, 0.0)
        eng.calibrate_bipolar_block(smp)
        ring_packets += pk
        out["eeg_cal"].append({"phase": "rest", "i": i + 1, "t": s.t})
    eng.finalize_calibration_bipolar()
    rest_feats = [ring_window(ring_packets[k * 300:(k + 1) * 300])[0] for k in range(RING_REST)]
    rest_feats = [f for f in rest_feats if f is not None]
    b_mean, b_std = compute_baseline(np.array([f[0] for f in rest_feats]), n_windows=len(rest_feats))
    score = lambda f: None if f is None else float(predict_window(f, b_mean, b_std)["prob_stress"])  # noqa: E731
    ring_rest = [score(f) for f in rest_feats]
    out["ring_cal"] += [{"phase": "rest", "i": k + 1, "p": v} for k, v in enumerate(ring_rest)]

    # Task (demand on) then recovery (demand off): EEG blocks and ring windows on the same clock.
    s.p.started_task = True
    for phase, n_eeg, n_ring, demand in (("task", N_TASK, RING_TASK, 1.0), ("recovery", N_REC, RING_REC, 0.0)):
        packets = []
        for i in range(n_eeg):
            smp, pk, loads = s.run(BLOCK_SEC, demand)
            packets += pk
            (eng.score_threshold_task_block if phase == "task" else eng.score_recovery_block)(smp)
            out["eeg_cal"].append({"phase": phase, "i": i + 1, "t": s.t, "loads": loads})
        need = n_ring * 300
        if len(packets) < need:                          # ring phase is slightly longer than the EEG one
            _, pk, _ = s.run((need - len(packets)) * STEP, demand)
            packets += pk
        for k in range(n_ring):
            out["ring_cal"].append({"phase": phase, "i": k + 1, "p": score(ring_window(packets[k * 300:(k + 1) * 300])[0])})
    thr = eng.finalize_personal_threshold()
    thr.update(eng.recovery_outcome())
    # One row per SCORED block, exactly the lists the decision used (untrusted
    # rest blocks -- e.g. the first, inside the electrode-settling window --
    # are left out by the backend, as live).
    out["eeg_cal"] = [{"phase": ph, "i": i + 1, "p": sc}
                      for ph, key in (("rest", "rest_scores"), ("task", "task_scores"), ("recovery", "recovery_scores"))
                      for i, sc in enumerate(thr.get(key) or [])]
    out["eeg_thr"] = thr
    rr = [r["p"] for r in out["ring_cal"] if r["phase"] == "rest" and r["p"] is not None]
    rt = [r["p"] for r in out["ring_cal"] if r["phase"] == "task" and r["p"] is not None]
    rc = [r["p"] for r in out["ring_cal"] if r["phase"] == "recovery" and r["p"] is not None]
    mid = midpoint_stress_threshold(rr, rt, recovery_scores=rc)
    out["ring_mid"] = mid
    ring_limit = mid["limit"] if mid["separated"] else STRESS_DECISION_THRESHOLD
    out["ring_limit"] = ring_limit

    # Live monitoring: one row per EEG block; the ring's latest 60 s window as the bridge sends it.
    history, ring_buf, k = [], [], 0
    for name, dur in LIVE:
        demand = 1.0 if name == "demand" else 0.0
        for _ in range(int(round(dur / BLOCK_SEC))):
            smp, pk, loads = s.run(BLOCK_SEC, demand)
            ring_buf = (ring_buf + pk)[-300:]
            res = eng.infer_bipolar_block(smp)
            history.append(res.get("calibrated_p") if not res.get("skipped") else None)
            rec = make_eeg_output_record(res, window_sec=BLOCK_SEC, model_version="bipolar_only_v1",
                                         feature_set="bipolar_only_v1", history_scores=history,
                                         high_load_p=eng.thresholds.high_load_p, capture_end_ts=1.8e9 + s.t)
            fe = make_eeg_fusion_record(rec)
            f, q = ring_window(ring_buf) if len(ring_buf) >= 300 else (None, None)
            pr = score(f)
            mods = {"eeg": {"prob": fe["prob_eeg"], "confidence": fe["conf_eeg"], "signal_quality": fe["signal_quality_eeg"],
                            "model_acc": fe["model_acc_eeg"], "threshold": fe["threshold_eeg"],
                            "threshold_scale": fe["threshold_scale_eeg"], "z": fe["z_logit_eeg"], "z_alert": fe["z_alert_eeg"]},
                    "hrv_gsr": {"prob": pr, "confidence": None if pr is None else max(pr, 1 - pr), "signal_quality": q,
                                "model_acc": 0.851, "threshold": ring_limit, "threshold_scale": PERSONAL_THRESHOLD_MIN_SEPARATION}}
            fu = fusion_engin.compute_realtime_multimodal_fusion(mods, scale="aligned")
            al = fu.get("aligned_probs", {})
            k += 1
            out["live"].append({
                "k": k, "t": s.t, "stage": name, "demand": demand, "loads": loads,
                "raw": rec.get("raw_task_state_score"), "score": rec.get("task_state_score"),
                "z": rec.get("task_state_z"), "z_alert": fe["z_alert_eeg"], "quality": rec.get("eeg_quality_score"),
                "skipped": bool(rec.get("skipped")), "mind": al.get("eeg"), "ring_p": pr, "body": al.get("hrv_gsr"),
                "fused": fu.get("fused_stress_probability"),
                "w_eeg": (fu.get("modality_weights") or {}).get("eeg"), "w_ring": (fu.get("modality_weights") or {}).get("hrv_gsr"),
            })
    return out


# ── Workbook ────────────────────────────────────────────────────────
def write_workbook(d, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule, FormulaRule

    F = "Arial"
    bold, norm = Font(name=F, bold=True), Font(name=F)
    blue = Font(name=F, color="0000FF")
    green = Font(name=F, color="008000")
    title = Font(name=F, bold=True, size=14)
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    thin = Border(bottom=Side(style="thin", color="B7C3D6"))
    wb = Workbook()

    def header(ws, row, cols, widths=None):
        for c, h in enumerate(cols, 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font, cell.fill, cell.border = bold, head_fill, thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    thr, mid = d["eeg_thr"], d["ring_mid"]

    # ── Summary (inputs first; the calibration sheets reference these cells) ──
    sm = wb.active
    sm.title = "Summary"
    sm["A1"] = f"Simulated session · profile '{d['profile']}' · real MindPulse backend pipeline"
    sm["A1"].font = title
    sm["A2"] = ("Blue = settings / expected values you can change · black = data from the backend run or formulas · "
                "green = links to other sheets. Every PASS/FAIL recomputes when inputs change.")
    sm["A2"].font = Font(name=F, italic=True, color="666666")
    sm.column_dimensions["A"].width = 46
    for c in "BCDE":
        sm.column_dimensions[c].width = 20

    def put(r, label, value, font=norm, fmt=None, note=None):
        sm.cell(row=r, column=1, value=label).font = norm
        c = sm.cell(row=r, column=2, value=value)
        c.font = font
        if fmt:
            c.number_format = fmt
        if note:
            c.comment = Comment(note, "simulator")
        return c

    r = 4
    sm.cell(row=r, column=1, value="Backend result: EEG calibration").font = bold
    rows_eeg = [
        ("Baseline (median rest log-odds)", thr.get("baseline_logit"), "0.000"),
        ("Spread (rest SD, log-odds)", thr.get("rest_sd_logit"), "0.000"),
        ("Sustained task level (log-odds)", thr.get("task_level_logit"), "0.000"),
        ("Limit (midpoint, raw log-odds)", thr.get("limit_logit"), "0.000"),
        ("Direction (+1 rises / -1 falls under load)", thr.get("personal_direction"), "0"),
        ("Separation (spreads vs rest)", thr.get("separation_sd"), "0.00"),
        ("Recovery reversal (spreads)", thr.get("reversal_sd"), "0.00"),
        ("Limit position on the z scale (z_alert)", thr.get("z_alert"), "0.000"),
        ("Calibration status", thr.get("calibration_status"), None),
        ("Task blocks skipped (ramp-up)", thr.get("n_task_skipped"), "0"),
    ]
    ref = {}
    for i, (lab, val, fmt) in enumerate(rows_eeg):
        put(r + 1 + i, lab, val, fmt=fmt)
        ref[lab] = f"Summary!$B${r + 1 + i}"
    r += len(rows_eeg) + 2
    sm.cell(row=r, column=1, value="Backend result: ring calibration").font = bold
    rows_ring = [
        ("Ring baseline (median rest)", mid["baseline"], "0.000"),
        ("Ring sustained task level", mid["task_level"], "0.000"),
        ("Ring separation needed", mid["needed_separation"], "0.000"),
        ("Ring limit in use", d["ring_limit"], "0.000"),
        ("Ring calibration status", mid["status"], None),
    ]
    for i, (lab, val, fmt) in enumerate(rows_ring):
        put(r + 1 + i, lab, val, fmt=fmt)
        ref[lab] = f"Summary!$B${r + 1 + i}"
    r += len(rows_ring) + 2

    # ── EEG calibration sheet ──
    ws = wb.create_sheet("EEG calibration")
    cols = ["Block", "Phase", "Model score (p)", "Log-odds", "Usable (|log-odds| ≤ 40)", "Plausible task #",
            "Rest log-odds (usable)", "Sustained task log-odds", "Recovery log-odds", "Oriented log-odds",
            "Fusion scale (limit = 0.50)", "Above limit"]
    header(ws, 1, cols, [8, 11, 16, 12, 14, 12, 14, 14, 14, 14, 16, 11])
    n = len(d["eeg_cal"])
    B, SD, DIR, LIM, ZA = ref["Baseline (median rest log-odds)"], ref["Spread (rest SD, log-odds)"], \
        ref["Direction (+1 rises / -1 falls under load)"], ref["Limit (midpoint, raw log-odds)"], \
        ref["Limit position on the z scale (z_alert)"]
    for i, row in enumerate(d["eeg_cal"]):
        x = i + 2
        ws.cell(row=x, column=1, value=i + 1)
        ws.cell(row=x, column=2, value=row["phase"])
        c = ws.cell(row=x, column=3, value=row["p"])
        c.number_format = "0.00E+00"
        ws.cell(row=x, column=4, value=f'=IFERROR(LN(C{x}/(1-C{x})),"")').number_format = "0.00"
        ws.cell(row=x, column=5, value=f'=IF(D{x}="",FALSE,ABS(D{x})<=40)')
        ws.cell(row=x, column=6, value=f'=IF(AND(B{x}="task",E{x}),COUNTIFS(B$2:B{x},"task",E$2:E{x},TRUE),"")')
        ws.cell(row=x, column=7, value=f'=IF(AND(B{x}="rest",E{x}),D{x},"")').number_format = "0.00"
        ws.cell(row=x, column=8, value=f'=IF(F{x}="","",IF(F{x}>Summary!$B$SKIP,D{x},""))').number_format = "0.00"
        ws.cell(row=x, column=9, value=f'=IF(AND(B{x}="recovery",E{x}),D{x},"")').number_format = "0.00"
        ws.cell(row=x, column=10, value=f'=IF(E{x},{B}+{DIR}*(D{x}-{B}),"")').number_format = "0.00"
        ws.cell(row=x, column=11, value=(f'=IF(E{x},1/(1+EXP(-MAX(-3,MIN(3,(J{x}-{B})/{SD}-{ZA})))),"")')).number_format = "0.00"
        ws.cell(row=x, column=12, value=f'=IF(K{x}="","",IF(K{x}>=0.5,1,0))')
        for cc in range(1, 13):
            ws.cell(row=x, column=cc).font = norm
    last = n + 1

    # ── Ring calibration sheet ──
    wr = wb.create_sheet("Ring calibration")
    header(wr, 1, ["Window (60 s)", "Phase", "Stress probability", "Phase index", "Rest", "Sustained task", "Recovery",
                   "Fusion scale (limit = 0.50)", "Above limit"], [12, 11, 16, 10, 10, 14, 10, 18, 11])
    RL = ref["Ring limit in use"]
    for i, row in enumerate(d["ring_cal"]):
        x = i + 2
        wr.cell(row=x, column=1, value=i + 1)
        wr.cell(row=x, column=2, value=row["phase"])
        wr.cell(row=x, column=3, value=row["p"]).number_format = "0.000"
        wr.cell(row=x, column=4, value=f'=COUNTIF(B$2:B{x},B{x})')
        wr.cell(row=x, column=5, value=f'=IF(AND(B{x}="rest",C{x}<>""),C{x},"")').number_format = "0.000"
        wr.cell(row=x, column=6, value=f'=IF(AND(B{x}="task",C{x}<>"",D{x}>Summary!$B$RSKIP),C{x},"")').number_format = "0.000"
        wr.cell(row=x, column=7, value=f'=IF(AND(B{x}="recovery",C{x}<>""),C{x},"")').number_format = "0.000"
        wr.cell(row=x, column=8, value=f'=IF(C{x}="","",1/(1+EXP(-MAX(-3,MIN(3,(C{x}-{RL})/{PERSONAL_THRESHOLD_MIN_SEPARATION})))))').number_format = "0.00"
        wr.cell(row=x, column=9, value=f'=IF(H{x}="","",IF(H{x}>=0.5,1,0))')
        for cc in range(1, 10):
            wr.cell(row=x, column=cc).font = norm
    rlast = len(d["ring_cal"]) + 1

    # ── Live monitoring sheet ──
    wl = wb.create_sheet("Live monitoring")
    lcols = ["Block", "Time (s)", "Scenario", "Demand", "Brain load", "Heart load", "Skin load",
             "EEG model score (raw)", "EEG score (oriented)", "EEG z", "z_alert",
             "Mind (fusion engine)", "Mind (check: from z)", "Ring stress prob", "Body (fusion engine)",
             "Fused (fusion engine)", "EEG weight", "Ring weight",
             "Mind ≥ limit", "Body ≥ limit", "Fused ≥ limit",
             "Mind alert (2 of 3)", "Body alert (2 of 3)", "Fused alert (2 of 3)",
             "Demand block #", "First fused alert in demand", "Back below after demand"]
    header(wl, 1, lcols, [7, 9, 10, 8, 9, 9, 9, 13, 13, 8, 8, 12, 12, 11, 12, 12, 9, 9, 9, 9, 9, 10, 10, 10, 10, 11, 11])
    for i, row in enumerate(d["live"]):
        x = i + 2
        vals = [row["k"], round(row["t"]), row["stage"], row["demand"], *[round(v, 3) for v in row["loads"]],
                row["raw"], row["score"], row["z"], row["z_alert"], row["mind"], None, row["ring_p"], row["body"],
                row["fused"], row["w_eeg"], row["w_ring"]]
        for cc, v in enumerate(vals, 1):
            c = wl.cell(row=x, column=cc, value=v)
            c.font = norm
        for cc in (8, 9):
            wl.cell(row=x, column=cc).number_format = "0.00E+00"
        for cc in (10, 11, 12, 14, 15, 16, 17, 18):
            wl.cell(row=x, column=cc).number_format = "0.00"
        wl.cell(row=x, column=13, value=f'=IF(J{x}="","",1/(1+EXP(-MAX(-3,MIN(3,J{x}-K{x})))))').number_format = "0.00"
        for cc, src in ((19, "L"), (20, "O"), (21, "P")):
            wl.cell(row=x, column=cc, value=f'=IF({src}{x}="","",IF({src}{x}>=0.5,1,0))')
        lo = max(2, x - 2)
        for cc, src in ((22, "S"), (23, "T"), (24, "U")):
            wl.cell(row=x, column=cc, value=f'=IF(COUNT({src}{lo}:{src}{x})=0,"",IF(SUM({src}{lo}:{src}{x})>=2,1,0))')
        wl.cell(row=x, column=25, value=f'=IF(C{x}="demand",A{x},"")')
        wl.cell(row=x, column=26, value=f'=IF(AND(C{x}="demand",X{x}=1),A{x},"")')
        wl.cell(row=x, column=27, value=f'=IF(AND(A{x}>MAX($Y$2:$Y$200),X{x}=0),A{x},"")')
    llast = len(d["live"]) + 1
    red = PatternFill("solid", fgColor="F8D7D3")
    wl.conditional_formatting.add(f"V2:X{llast}", CellIsRule(operator="equal", formula=["1"], fill=red))
    wl.conditional_formatting.add(f"A2:D{llast}", FormulaRule(formula=['$C2="demand"'], fill=PatternFill("solid", fgColor="FFF4D6")))

    # ── Summary: settings, Excel recomputation, checks ──
    sm.cell(row=r, column=1, value="Settings (the backend's constants)").font = bold
    settings = [("EEG task ramp-up blocks skipped (max)", TASK_RAMP_SKIP_BLOCKS, "sam40/decision.py TASK_RAMP_SKIP_BLOCKS"),
                ("EEG separation needed (spreads)", MIN_TASK_RESPONSE_SD, "sam40/decision.py MIN_TASK_RESPONSE_SD"),
                ("EEG recovery reversal needed (spreads)", MIN_RECOVERY_REVERSAL_SD, "sam40/decision.py MIN_RECOVERY_REVERSAL_SD"),
                ("EEG spread floor (log-odds)", 0.55, "sam40/decision.py PERSONAL_SIGMA_LOGIT"),
                ("Ring task ramp-up windows skipped (max)", TASK_RAMP_SKIP_WINDOWS, "physiological/predictor.py TASK_RAMP_SKIP_WINDOWS"),
                ("Ring minimum separation", PERSONAL_THRESHOLD_MIN_SEPARATION, "physiological/predictor.py PERSONAL_THRESHOLD_MIN_SEPARATION")]
    srow = {}
    for i, (lab, val, note) in enumerate(settings):
        put(r + 1 + i, lab, val, font=blue, note=f"Source: backend constant {note}")
        srow[lab] = r + 1 + i
    r += len(settings) + 1
    # effective skips (ramp_skip keeps at least half of a short task)
    put(r, "EEG blocks skipped (effective)", f'=MIN(B{srow["EEG task ramp-up blocks skipped (max)"]},INT(COUNTIFS(\'EEG calibration\'!B2:B{last},"task",\'EEG calibration\'!E2:E{last},TRUE)/2))', fmt="0")
    skip_row = r
    put(r + 1, "Ring windows skipped (effective)", f'=MIN(B{srow["Ring task ramp-up windows skipped (max)"]},INT(COUNTIFS(\'Ring calibration\'!B2:B{rlast},"task",\'Ring calibration\'!C2:C{rlast},"<>")/2))', fmt="0")
    rskip_row = r + 1
    r += 3

    # patch the placeholder references to the effective-skip cells
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=8, max_col=8):
        row[0].value = row[0].value.replace("$B$SKIP", f"$B${skip_row}")
    for row in wr.iter_rows(min_row=2, max_row=rlast, min_col=6, max_col=6):
        row[0].value = row[0].value.replace("$B$RSKIP", f"$B${rskip_row}")

    E = "'EEG calibration'!"
    Rg = "'Ring calibration'!"
    L = "'Live monitoring'!"
    sm.cell(row=r, column=1, value="Excel's own recomputation (from the per-block data)").font = bold
    sm.cell(row=r, column=2, value="Excel").font = bold
    sm.cell(row=r, column=3, value="Backend").font = bold
    sm.cell(row=r, column=4, value="Match").font = bold
    sd_floor = f'B{srow["EEG spread floor (log-odds)"]}'
    rec = [
        ("Baseline = median of usable rest log-odds", f"=MEDIAN({E}G2:G{last})", B, 0.001),
        ("Spread = max(floor, SD of rest) with ≥5 blocks", f"=IF(COUNT({E}G2:G{last})>=5,MAX({sd_floor},STDEV({E}G2:G{last})),{sd_floor})", SD, 0.001),
        ("Sustained task level = mean after ramp-up", f"=AVERAGE({E}H2:H{last})", ref["Sustained task level (log-odds)"], 0.001),
        ("Limit = midpoint of baseline and task", "=(B{a}+B{c})/2", LIM, 0.001),
        ("Direction = sign(task - baseline)", "=IF(B{c}>=B{a},1,-1)", DIR, 0),
        ("Separation = (task - baseline) / spread", "=(B{c}-B{a})/B{b}", ref["Separation (spreads vs rest)"], 0.001),
        ("Recovery reversal = dir x (task - median recovery) / spread", f"=B{{e}}*(B{{c}}-MEDIAN({E}I2:I{last}))/B{{b}}", ref["Recovery reversal (spreads)"], 0.001),
    ]
    base_row = r + 1
    rr_ = {"a": base_row, "b": base_row + 1, "c": base_row + 2, "e": base_row + 4}
    for i, (lab, f_excel, backend, tol) in enumerate(rec):
        x = base_row + i
        sm.cell(row=x, column=1, value=lab).font = norm
        sm.cell(row=x, column=2, value=f_excel.format(**rr_)).number_format = "0.000"
        c = sm.cell(row=x, column=3, value=f"={backend}")
        c.font, c.number_format = green, "0.000"
        sm.cell(row=x, column=4, value=f'=IF(ABS(B{x}-C{x})<={tol if tol else 0.0001},"PASS","FAIL")')
    x = base_row + len(rec)
    sep_need, rev_need = f'B{srow["EEG separation needed (spreads)"]}', f'B{srow["EEG recovery reversal needed (spreads)"]}'
    sm.cell(row=x, column=1, value="Decision = responsive if |separation| ≥ 1 and reversal ≥ 0.5").font = norm
    sm.cell(row=x, column=2, value=f'=IF(ABS(B{base_row + 5})>={sep_need},IF(B{base_row + 6}>={rev_need},"ok","drift_suspected"),"not_separated")')
    sm.cell(row=x, column=3, value=f"={ref['Calibration status']}").font = green
    sm.cell(row=x, column=4, value=f'=IF(B{x}=C{x},"PASS","FAIL")')
    x += 1
    rb, rt_, rn = x, x + 1, x + 2
    sm.cell(row=rb, column=1, value="Ring baseline = median of rest windows").font = norm
    sm.cell(row=rb, column=2, value=f"=MEDIAN({Rg}E2:E{rlast})").number_format = "0.000"
    sm.cell(row=rb, column=3, value=f"={ref['Ring baseline (median rest)']}").font = green
    sm.cell(row=rb, column=4, value=f'=IF(ABS(B{rb}-C{rb})<=0.001,"PASS","FAIL")')
    sm.cell(row=rt_, column=1, value="Ring sustained task = mean after ramp-up").font = norm
    sm.cell(row=rt_, column=2, value=f"=AVERAGE({Rg}F2:F{rlast})").number_format = "0.000"
    sm.cell(row=rt_, column=3, value=f"={ref['Ring sustained task level']}").font = green
    sm.cell(row=rt_, column=4, value=f'=IF(ABS(B{rt_}-C{rt_})<=0.001,"PASS","FAIL")')
    ring_min = f'B{srow["Ring minimum separation"]}'
    sm.cell(row=rn, column=1, value="Ring decision (separation ≥ max(SD, 0.10), recovery comes back)").font = norm
    sm.cell(row=rn, column=2, value=(f'=IF(B{rt_}-B{rb}>=MAX(STDEV({Rg}E2:E{rlast}),{ring_min}),'
                                    f'IF(COUNT({Rg}G2:G{rlast})=0,"ok",IF(B{rt_}-MEDIAN({Rg}G2:G{rlast})>=0.5*MAX(STDEV({Rg}E2:E{rlast}),{ring_min}),"ok","drift_suspected")),"not_separated")'))
    sm.cell(row=rn, column=3, value=f"={ref['Ring calibration status']}").font = green
    sm.cell(row=rn, column=4, value=f'=IF(B{rn}=C{rn},"PASS","FAIL")')
    fx = rn + 1
    sm.cell(row=fx, column=1, value="Mind value recomputed from z and z_alert (max difference)").font = norm
    sm.cell(row=fx, column=2, value=f"=SUMPRODUCT(MAX(ABS(N(+{L}L2:L{llast})-N(+{L}M2:M{llast}))))").number_format = "0.0000"
    sm.cell(row=fx, column=4, value=f'=IF(B{fx}<=0.001,"PASS","FAIL")')

    # Behaviour checks: expected (blue, editable) vs actual (formula)
    r = fx + 2
    sm.cell(row=r, column=1, value="Behaviour checks").font = bold
    for c, h in enumerate(["Actual", "Expected", "Result"], 2):
        sm.cell(row=r, column=c, value=h).font = bold
    exp_status = {"typical": "ok", "inverted": "ok", "drift": "drift_suspected"}[d["profile"]]
    exp_dir = {"typical": 1, "inverted": -1, "drift": 1}[d["profile"]]
    good = d["profile"] != "drift"
    dem = f'COUNTIF({L}C2:C{llast},"demand")'
    calm = f'COUNTIF({L}C2:C{llast},"calm")'
    checks = [
        ("EEG calibration status", f"={ref['Calibration status']}", exp_status, "eq"),
        ("EEG direction", f"={DIR}", exp_dir, "eq"),
        ("Ring calibration status", f"={ref['Ring calibration status']}", "ok" if good else "not_separated", "eq"),
        ("Fused alert rate during demand", f'=COUNTIFS({L}C2:C{llast},"demand",{L}X2:X{llast},1)/{dem}', 0.7 if good else 0.0, "ge"),
        ("Fused false-alert rate while calm", f'=COUNTIFS({L}C2:C{llast},"calm",{L}X2:X{llast},1)/{calm}', 0.10, "le"),
        ("Mind alert rate during demand", f'=COUNTIFS({L}C2:C{llast},"demand",{L}V2:V{llast},1)/{dem}', 0.6 if good else 0.0, "ge"),
        ("Mind false-alert rate while calm", f'=COUNTIFS({L}C2:C{llast},"calm",{L}V2:V{llast},1)/{calm}', 0.15 if good else 1.0, "le"),
        ("Body alert rate during demand", f'=COUNTIFS({L}C2:C{llast},"demand",{L}W2:W{llast},1)/{dem}', 0.6 if good else 0.0, "ge"),
        ("Body false-alert rate while calm", f'=COUNTIFS({L}C2:C{llast},"calm",{L}W2:W{llast},1)/{calm}', 0.15, "le"),
        ("Blocks from demand start to first fused alert", f'=IF(COUNT({L}Z2:Z{llast})=0,"none",MIN({L}Z2:Z{llast})-MIN({L}Y2:Y{llast}))', 6 if good else "none", "le" if good else "eq"),
        ("Blocks from demand end until fused is back below", f'=IF(COUNT({L}AA2:AA{llast})=0,"never",MIN({L}AA2:AA{llast})-MAX({L}Y2:Y{llast}))', 6, "le"),
    ]
    for i, (lab, actual, expected, op) in enumerate(checks):
        x = r + 1 + i
        sm.cell(row=x, column=1, value=lab).font = norm
        a = sm.cell(row=x, column=2, value=actual)
        e = sm.cell(row=x, column=3, value=expected)
        e.font = blue
        if "rate" in lab:
            a.number_format = e.number_format = "0%"
        cond = {"eq": f"B{x}=C{x}", "ge": f"B{x}>=C{x}", "le": f"B{x}<=C{x}"}[op]
        if op != "eq":
            cond = f'AND(ISNUMBER(B{x}),{cond})' if expected != "none" else cond
        sm.cell(row=x, column=4, value=f'=IF({cond},"PASS","FAIL")')
        e.comment = Comment("Expected value for this profile (set by the simulator author) -- edit to test other criteria.", "simulator")
    x = r + 1 + len(checks) + 1
    sm.cell(row=x, column=1, value="Overall").font = bold
    sm.cell(row=x, column=4, value=f'=IF(COUNTIF(D{base_row}:D{x - 1},"FAIL")=0,"ALL PASS",COUNTIF(D{base_row}:D{x - 1},"FAIL")&" FAIL")').font = bold
    okfill, badfill = PatternFill("solid", fgColor="D9F2DD"), PatternFill("solid", fgColor="F8D7D3")
    sm.conditional_formatting.add(f"D{base_row}:D{x}", FormulaRule(formula=[f'ISNUMBER(SEARCH("PASS",D{base_row}))'], fill=okfill))
    sm.conditional_formatting.add(f"D{base_row}:D{x}", FormulaRule(formula=[f'ISNUMBER(SEARCH("FAIL",D{base_row}))'], fill=badfill))
    for row in sm.iter_rows(min_row=4, max_row=x):
        for c in row:
            if c.font == Font():
                c.font = norm
    sm.cell(row=x + 2, column=1, value=("Protocol: 2 min settle (discarded) -> rest 10 EEG blocks / 4 ring windows -> task 19 / 8 "
                                        "(demand on) -> recovery 7 / 3 (demand off); then live: calm 4 min -> demand 5 min -> calm 5 min."
                                        )).font = Font(name=F, italic=True, color="666666")
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--profile", default="typical", choices=("typical", "inverted", "drift"))
    args = ap.parse_args()
    S.RNG = np.random.default_rng(11)
    d = simulate(args.profile)
    out = REPO / "backend" / "simulation" / "output" / f"simulated_session_{args.profile}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(d, out)
    print(out)


if __name__ == "__main__":
    main()
